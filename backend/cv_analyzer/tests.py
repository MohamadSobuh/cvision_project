import io
from unittest.mock import patch

from django.test import SimpleTestCase
from docx import Document

from .ai import services
from .ai.services import analyze_cv_gaps, extract_text_from_docx, normalize_field


class CVAnalyzerTests(SimpleTestCase):
    def test_backend_development_label_is_normalized(self):
        self.assertEqual(normalize_field("Back-end Development"), "backend")
        self.assertEqual(normalize_field("back_end_development"), "backend")
        self.assertEqual(normalize_field("BACK END DEVELOPMENT"), "backend")

    def test_docx_extraction_includes_table_cells(self):
        document = Document()
        document.add_paragraph("Backend engineer")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Frameworks"
        table.cell(0, 1).text = "Django, Flask, FastAPI"

        output = io.BytesIO()
        document.save(output)

        extracted_text = extract_text_from_docx(output.getvalue())

        self.assertIn("Backend engineer", extracted_text)
        self.assertIn("Django, Flask, FastAPI", extracted_text)

    def test_mysql_does_not_count_as_standalone_sql(self):
        document = Document()
        document.add_paragraph("MySQL")
        output = io.BytesIO()
        document.save(output)
        output.seek(0)

        strengths, weaknesses, _ = analyze_cv_gaps(
            output, "Back-end Development"
        )

        strength_names = {item["skill"] for item in strengths}
        weakness_names = {item["skill"] for item in weaknesses}
        self.assertIn("mysql", strength_names)
        self.assertNotIn("sql", strength_names)
        self.assertIn("sql", weakness_names)

    def test_bert_ner_entities_are_used_for_skill_detection(self):
        class FakeNerPipeline:
            def __call__(self, chunks):
                return [[{
                    "word": "Django",
                    "score": 0.98,
                    "entity_group": "MISC",
                }]]

        document = Document()
        document.add_paragraph("Backend engineer experienced in web services")
        output = io.BytesIO()
        document.save(output)
        output.seek(0)

        with patch.object(services, "_ner_pipeline", FakeNerPipeline()):
            strengths, _, _ = analyze_cv_gaps(
                output, "Back-end Development"
            )

        django_strength = next(
            item for item in strengths if item["skill"] == "django"
        )
        self.assertEqual(django_strength["source"], "bert_ner")
        self.assertEqual(django_strength["matched_text"], "Django")
