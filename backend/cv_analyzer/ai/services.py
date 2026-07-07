
import io
import logging
import re
import threading

from django.conf import settings
from docx import Document
import pandas as pd
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.naive_bayes import MultinomialNB
import pdfplumber


logger = logging.getLogger(__name__)

_ner_pipeline = None
_ner_model_name = None
_ner_initialization_attempted = False
_ner_initialization_lock = threading.Lock()

# ==========================================
# 1. قاموس المهارات الشامل والمتاح في النظام (تعديل التسميات الموحدة)
# ==========================================
SKILL_DICTIONARY = {
    # === Frontend Skills ===
    'html': [r'\bhtml\b', r'\bhtml5\b'],
    'css': [r'\bcss\b', r'\bcss3\b'],
    'javascript': [r'\bjavascript\b', r'\bjs\b'],
    'react': [r'\breact\b', r'\breact\.js\b', r'\breactjs\b'],
    'tailwind': [r'\btailwind\b', r'\btailwindcss\b', r'\btailwind css\b'],
    'typescript': [r'\btypescript\b', r'\bts\b'],
    'next.js': [r'\bnext\.js\b', r'\bnextjs\b'],
    'redux': [r'\bredux\b', r'\bredux toolkit\b', r'\brtk\b'],
    'sass': [r'\bsass\b', r'\bscss\b'],
    'bootstrap': [r'\bbootstrap\b', r'\bbootstrap[4-5]\b'],
    'responsive': [r'\bresponsive web design\b', r'\bresponsive design\b', r'\brwd\b', r'\bresponsive\b'],

    # === Backend & Databases ===
    'python': [r'\bpython\b', r'\bpython3\b', r'\bpy\b'],
    'django': [r'\bdjango\b', r'\bdjango rest framework\b', r'\bdrf\b'],
    'flask': [r'\bflask\b'],
    'fastapi': [r'\bfastapi\b', r'\bfast api\b'],
    'rest_api': [r'\brest api\b', r'\brestful api\b', r'\brest apis\b', r'\bapis\b'],
    'sql': [r'\bsql\b', r'\bsql queries\b'],
    'postgresql': [r'\bpostgresql\b', r'\bpostgres\b', r'\bpostgres sql\b'],
    'mysql': [r'\bmysql\b', r'\bmy sql\b'],
    'mongodb': [r'\bmongodb\b', r'\bmongo db\b', r'\bmongo\b'],
    'docker': [r'\bdocker\b', r'\bdocker containers\b'],
    'node.js': [r'\bnode\.js\b', r'\bnodejs\b', r'\bnode\b'],
    'express': [r'\bexpress\b', r'\bexpress\.js\b', r'\bexpressjs\b'],
    'aws': [r'\baws\b', r'\bamazon web services\b'],

    # === Data Science & Analytics ===
    'pandas': [r'\bpandas\b'],
    'numpy': [r'\bnumpy\b'],
    'matplotlib': [r'\bmatplotlib\b'],
    'seaborn': [r'\bseaborn\b'],
    'statistics': [r'\bstatistics\b', r'\bstatistical analysis\b'],
    'tableau': [r'\btableau\b'],
    'power_bi': [r'\bpower bi\b', r'\bpowerbi\b', r'\bpbi\b'],
    'excel': [r'\bexcel\b', r'\bms excel\b', r'\bmicrosoft excel\b'],
    'data_visualization': [r'\bdata visualization\b', r'\bvisualization\b', r'\bdata viz\b'],

    # === Artificial Intelligence & Machine Learning ===
    'artificial_intelligence': [r'\bartificial intelligence\b', r'\bai\b'],
    'machine_learning': [r'\bmachine learning\b', r'\bml\b'],
    'deep_learning': [r'\bdeep learning\b', r'\bdl\b'],
    'nlp': [r'\bnlp\b', r'\bnatural language processing\b'],
    'computer_vision': [r'\bcomputer vision\b', r'\bcv\b'],
    'opencv': [r'\bopencv\b', r'\bopen cv\b'],
    'bert': [r'\bbert\b', r'\bllm bert\b'],
    'rag': [r'\brag\b', r'\bretrieval augmented generation\b'],
    'agentic_ai': [r'\bagentic ai\b', r'\bai agents\b', r'\bai agent\b'],
    'tensorflow': [r'\btensorflow\b', r'\btf\b'],
    'pytorch': [r'\bpytorch\b', r'\btorch\b'],
    'scikit_learn': [r'\bscikit-learn\b', r'\bscikit learn\b', r'\bsklearn\b']
}

# توليد القائمة تلقائياً لضمان سلامة الـ Vectorizer والموديل
ALL_SYSTEM_SKILLS = list(SKILL_DICTIONARY.keys())

SKILL_DISPLAY_NAMES = {
    'html': 'HTML5',
    'css': 'CSS3',
    'javascript': 'JavaScript',
    'react': 'React.js',
    'tailwind': 'Tailwind CSS',
    'typescript': 'TypeScript',
    'next.js': 'Next.js',
    'redux': 'Redux Toolkit',
    'sass': 'Sass / SCSS',
    'bootstrap': 'Bootstrap',
    'responsive': 'Responsive Web Design',
    'python': 'Python',
    'django': 'Django (DRF)',
    'flask': 'Flask',
    'fastapi': 'FastAPI',
    'rest_api': 'RESTful API',
    'sql': 'SQL',
    'postgresql': 'PostgreSQL',
    'mysql': 'MySQL',
    'mongodb': 'MongoDB',
    'docker': 'Docker',
    'node.js': 'Node.js',
    'express': 'Express.js',
    'aws': 'AWS Cloud',
    'pandas': 'Pandas',
    'numpy': 'NumPy',
    'matplotlib': 'Matplotlib',
    'seaborn': 'Seaborn',
    'statistics': 'Statistical Analysis',
    'tableau': 'Tableau',
    'power_bi': 'Power BI',
    'excel': 'Microsoft Excel',
    'data_visualization': 'Data Visualization',
    'artificial_intelligence': 'Artificial Intelligence (AI)',
    'machine_learning': 'Machine Learning (ML)',
    'deep_learning': 'Deep Learning (DL)',
    'nlp': 'Natural Language Processing (NLP)',
    'computer_vision': 'Computer Vision',
    'opencv': 'OpenCV Library',
    'bert': 'BERT Model',
    'rag': 'RAG Architecture',
    'agentic_ai': 'Agentic AI / AI Agents',
    'tensorflow': 'TensorFlow',
    'pytorch': 'PyTorch',
    'scikit_learn': 'Scikit-Learn'
}

# ==========================================
# 2. تجهيز بيانات تدريب الموديل (Training Dataset)
# ==========================================
training_data = {
    "text": [
        "html css javascript react tailwind css typescript next.js redux sass bootstrap responsive frontend ui",
        "html5 css3 js reactjs tailwindcss sass web design frontend",
        "python django flask fastapi rest api sql postgresql mysql mongodb backend node.js express aws server",
        "backend engineer python django sql apis restful nodejs databases docker backend architecture",
        "python artificial intelligence machine learning deep learning nlp computer vision opencv bert rag agentic ai tensorflow pytorch scikit-learn neural networks",
        "ai engineer intelligent agents large language models llms natural language processing prompt engineering computer vision object detection",
        "python pandas numpy matplotlib seaborn sql statistics data visualization",
        "data analyst excel tableau power bi sql python pandas visualization",
    ],
    "label": [
        "frontend",
        "frontend",
        "backend",
        "backend",
        "artificial_intelligence",
        "artificial_intelligence",
        "data_analysis",
        "data_analysis",
    ],
}

df_train = pd.DataFrame(training_data)

# ==========================================
# 3. خطوة تدريب الموديل (Model Training)
# ==========================================
logger.info("[CV Analyzer] Training the Naive Bayes classifier...")
vectorizer = TfidfVectorizer()
X_train = vectorizer.fit_transform(df_train["text"])
y_train = df_train["label"]

ml_model = MultinomialNB(alpha=1.0)
ml_model.fit(X_train, y_train)


def initialize_bert_ner():
    """
    Load the Hugging Face BERT NER pipeline once per Django process.

    AppConfig.ready() calls this during startup. If loading fails, the current
    regex-based detector remains available as a safe fallback.
    """
    global _ner_pipeline
    global _ner_model_name
    global _ner_initialization_attempted

    if not getattr(settings, "CV_ANALYZER_LOAD_BERT_NER", True):
        logger.info("[CV Analyzer] BERT-NER loading is disabled.")
        return False

    if _ner_pipeline is not None:
        return True

    with _ner_initialization_lock:
        if _ner_pipeline is not None:
            return True
        if _ner_initialization_attempted:
            return False

        _ner_initialization_attempted = True
        model_name = getattr(
            settings,
            "CV_ANALYZER_BERT_NER_MODEL",
            "dslim/bert-base-NER",
        )
        device = getattr(settings, "CV_ANALYZER_BERT_NER_DEVICE", -1)

        try:
            from transformers import pipeline

            logger.info(
                "[CV Analyzer] Loading BERT-NER model '%s'...",
                model_name,
            )
            _ner_pipeline = pipeline(
                task="ner",
                model=model_name,
                tokenizer=model_name,
                aggregation_strategy="simple",
                device=device,
            )
            _ner_model_name = model_name
            logger.info("[CV Analyzer] BERT-NER model loaded successfully.")
            return True
        except Exception:
            logger.exception(
                "[CV Analyzer] Could not load BERT-NER. "
                "CV analysis will continue with regex matching."
            )
            return False


def _split_ner_text(text, max_chars=1200, max_chunks=8):
    """Split long CV text into chunks that stay below BERT's token limit."""
    chunks = []
    remaining = text.strip()

    while remaining and len(chunks) < max_chunks:
        if len(remaining) <= max_chars:
            chunks.append(remaining)
            break

        split_at = remaining.rfind(" ", 0, max_chars)
        if split_at <= 0:
            split_at = max_chars

        chunks.append(remaining[:split_at].strip())
        remaining = remaining[split_at:].strip()

    return [chunk for chunk in chunks if chunk]


def extract_named_entities(text):
    """
    Return deduplicated BERT-NER entities or an empty list if unavailable.
    """
    if _ner_pipeline is None or not text:
        return []

    chunks = _split_ner_text(text)
    if not chunks:
        return []

    try:
        chunk_results = _ner_pipeline(chunks)
    except Exception:
        logger.exception(
            "[CV Analyzer] BERT-NER inference failed; using regex fallback."
        )
        return []

    if chunk_results and isinstance(chunk_results[0], dict):
        chunk_results = [chunk_results]

    min_confidence = getattr(
        settings,
        "CV_ANALYZER_NER_MIN_CONFIDENCE",
        0.85,
    )
    entities_by_name = {}

    for entities in chunk_results:
        for entity in entities:
            word = entity.get("word", "").replace("##", "").strip()
            score = float(entity.get("score", 0.0))
            if not word or score < min_confidence:
                continue

            normalized_word = word.casefold()
            previous = entities_by_name.get(normalized_word)
            if previous is None or score > previous["score"]:
                entities_by_name[normalized_word] = {
                    "word": word,
                    "score": score,
                    "label": entity.get("entity_group")
                    or entity.get("entity", ""),
                }

    return list(entities_by_name.values())


def _find_skill_in_entities(patterns, named_entities):
    for entity in named_entities:
        entity_text = entity["word"].lower()
        if any(re.search(pattern, entity_text) for pattern in patterns):
            return entity
    return None
logger.info("[CV Analyzer] Naive Bayes classifier trained successfully.")

# ==========================================
# 5. خريطة تطبيع أسماء المجالات
# ==========================================
FIELD_LABEL_MAP = {
    "frontend": "frontend",
    "front end": "frontend",
    "front-end": "frontend",
    "frontend development": "frontend",
    "front-end development": "frontend",
    "front end development": "frontend",
    "backend": "backend",
    "back end": "backend",
    "back-end": "backend",
    "back-end development": "backend",
    "back end development": "backend",
    "artificial intelligence": "artificial_intelligence",
    "data analysis": "data_analysis",
    "data analytics": "data_analysis",
    "data analyze": "data_analysis",
    "data analyzes": "data_analysis",
    "data-analytics": "data_analysis",
    "data_analysis": "data_analysis",
}


def normalize_field(target_field: str) -> str:
    key = target_field.lower().strip()
    key = re.sub(r"[\s_-]+", " ", key)
    if key in FIELD_LABEL_MAP:
        return FIELD_LABEL_MAP[key]
    field_vector = vectorizer.transform([key])
    return ml_model.predict(field_vector)[0]


# ==========================================
# 6. تحديث مصفوفة الـ Hardcoded لتطابق القاموس الجديد تماماً 
# ==========================================
HARDCODED_FIELD_SKILLS = {
    "frontend": ["html", "css", "javascript", "react", "tailwind", "typescript", "next.js", "redux", "sass", "bootstrap", "responsive"],
    "backend": ["python", "django", "flask", "fastapi", "rest_api", "sql", "postgresql", "mysql", "mongodb", "docker", "node.js", "express", "aws"],
    "artificial_intelligence": ["python", "artificial_intelligence", "machine_learning", "deep_learning", "nlp", "computer_vision", "opencv", "bert", "rag", "agentic_ai", "tensorflow", "pytorch", "scikit_learn"],
    "data_analysis": ["python", "pandas", "numpy", "sql", "statistics", "data_visualization", "matplotlib", "seaborn", "tableau", "power_bi", "excel"],
}


def extract_skills_by_model_prediction(target_field):
    label = normalize_field(target_field)
    if label in HARDCODED_FIELD_SKILLS:
        return HARDCODED_FIELD_SKILLS[label]

    learned_skills = []
    for skill in ALL_SYSTEM_SKILLS:
        skill_vector = vectorizer.transform([skill])
        predicted_category = ml_model.predict(skill_vector)[0]
        if predicted_category == label:
            learned_skills.append(skill)
    return learned_skills


# ==========================================
# 6. دالة تنظيف النص
# ==========================================
def cleanText(text):
    cleanedText = text.replace("■", "")
    cleanedText = cleanedText.replace("\n", " ")
    cleanedText = cleanedText.replace("/n", " ")
    cleanedText = cleanedText.replace("\\n", " ")
    cleanedText = re.sub(r"[^A-Za-z0-9\s\+#\.]", " ", cleanedText)
    cleanedText = re.sub(r"^n\s+", "", cleanedText)
    cleanedText = re.sub(r"\s+n\s+", " ", cleanedText)
    cleanedText = re.sub(r"\s+", " ", cleanedText).strip()
    return cleanedText


# ==========================================
# 7. دوال استخراج النصوص من الذاكرة
# ==========================================
def extract_text_from_docx(raw_bytes):
    document = Document(io.BytesIO(raw_bytes))
    text_blocks = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            text_blocks.extend(
                cell.text for cell in row.cells if cell.text
            )
    return "\n".join(text_blocks)


def extract_text_from_uploaded_file(file_obj):
    raw_bytes = file_obj.read()
    all_text = ""
    if raw_bytes[:4] == b"%PDF":
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            for page in pdf.pages:
                text_page = page.extract_text()
                if text_page:
                    all_text += text_page + "\n"
    elif raw_bytes[:2] == b"PK":
        all_text = extract_text_from_docx(raw_bytes)
    else:
        raise ValueError("Unsupported file format for text extraction.")
    return cleanText(all_text)


# ==========================================
# 8. الدالة الرئيسية لتحليل الفجوات والمطابقة (Gap Analysis) - النسخة المصححة
# ==========================================
def analyze_cv_gaps(file_obj, target_field):
    logger.debug("Analyzing CV for target field: %s", target_field)
    
    # 1. استخراج المهارات المطلوبة للمجال المطلوب
    required_skills = extract_skills_by_model_prediction(target_field)
    if not required_skills:
        return [], [], []

    # 2. قراءة وتنظيف النص من الملف المرفوع
    cleaned_text = extract_text_from_uploaded_file(file_obj)
    text_lower = cleaned_text.lower()
    named_entities = extract_named_entities(cleaned_text)

    strengths = []
    found_skills = set()
    required_skills_set = set(required_skills)

    # 4. فحص مهارات المجال المحدد فقط بحثاً عن نقاط القوة
    for skill_name, patterns in SKILL_DICTIONARY.items():
        if skill_name not in required_skills_set:
            continue

        entity_match = _find_skill_in_entities(patterns, named_entities)
        if entity_match:
            strengths.append({
                "skill": skill_name,
                "matched_text": entity_match["word"],
                "confidence": round(entity_match["score"], 4),
                "source": "bert_ner",
            })
            found_skills.add(skill_name)
            continue

        for pattern in patterns:
            match = re.search(pattern, text_lower)
            if match:

                display_name = SKILL_DISPLAY_NAMES.get(skill_name, skill_name.title())
                
                # حساب الـ Confidence
                if match and len(match.group(0)) > len(skill_name):
                    confidence = 0.95
                else:
                    confidence = 0.90
                    
                strengths.append({
                    "skill": skill_name,
                    "matched_text": display_name,
                    "confidence": confidence,
                    "source": "regex",
                })    
                found_skills.add(skill_name)
                break  # نكتفي بمطابقة واحدة للمهارة لمنع التكرار

    # 5. استخراج نقاط الضعف (المهارات المطلوبة وغير الموجودة في found_skills)
    weaknesses = []
    for required_skill in required_skills:
        if required_skill not in found_skills:
            weaknesses.append({
                "skill": required_skill,
                "reason": f"Missing core skill required for the {target_field} track.",
            })   

    logger.debug("Detected strengths: %s", strengths)
    logger.debug("Detected weaknesses: %s", weaknesses)

    # أصبحت المصفوفات مرتبة ومحمية من التكرار عبر البنية السابقة، نرجعها مباشرة دون كراش
    return strengths, weaknesses, required_skills
